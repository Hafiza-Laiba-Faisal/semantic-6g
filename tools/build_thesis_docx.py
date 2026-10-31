#!/usr/bin/env python3
"""Build a Word document from a Markdown source file.

The Markdown file is the single source of truth. This converter handles the
subset of Markdown used by the thesis and figure documents: headings,
paragraphs, bullet/numbered lists, pipe tables, blockquotes, images,
bold/italic/code spans, and explicit page-break markers.

Usage:
    python tools/build_thesis_docx.py
    python tools/build_thesis_docx.py --md fyp_md_files/FIGURES_DOCUMENT.md \
        --out FIGURES_DOCUMENT.docx --header "Thesis Figures"
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "fyp_md_files" / "THESIS_MASTER_DOCUMENT.md"
OUT_PATH = ROOT / "THESIS_MASTER_DOCUMENT.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT_NAVY = "DCE6F1"

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
IMAGE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
PAGEBREAK = "<!-- pagebreak -->"
TOC_MARKER = "<!-- toc -->"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_page_number_footer(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_field(p, "PAGE", "1")
    run = p.add_run(" of ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_field(p, "NUMPAGES", "1")


def add_header(section, text: str) -> None:
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    run.italic = True


def add_toc(doc: Document) -> None:
    # Deliberately a plain paragraph, not a Heading, so the TOC does not list
    # itself.
    heading = doc.add_paragraph()
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.space_after = Pt(12)
    run = heading.add_run("Table of Contents")
    run.font.name = "Cambria"
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = NAVY
    p = doc.add_paragraph()
    add_field(
        p,
        'TOC \\o "1-3" \\h \\z \\u',
        "Right-click here and choose Update Field to build the table of contents.",
    )
    doc.add_page_break()


def enable_update_fields(doc: Document) -> None:
    """Ask Word to refresh the TOC and page fields when the file is opened."""
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_inline(paragraph, text: str, size: float | None = None,
               bold: bool = False, italic: bool = False,
               color: RGBColor | None = None) -> None:
    """Add text to a paragraph, honouring **bold**, *italic*, and `code`."""
    for token in INLINE.split(text):
        if not token:
            continue
        is_bold = bold
        is_italic = italic
        is_code = False
        content = token
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            content = token[2:-2]
            is_bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            content = token[1:-1]
            is_italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            content = token[1:-1]
            is_code = True
        run = paragraph.add_run(content)
        run.bold = is_bold
        run.italic = is_italic
        if is_code:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 1.5)
        elif size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def add_paragraph(doc: Document, text: str, *, center: bool = False,
                  size: float | None = None, bold: bool = False,
                  italic: bool = False, color: RGBColor | None = None,
                  space_after: float = 6) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, text, size=size, bold=bold, italic=italic, color=color)


def add_table(doc: Document, rows: list[list[str]], *, center: bool = False) -> None:
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.autofit = True
    if center:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows):
        for c_idx in range(n_cols):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            value = row[c_idx] if c_idx < len(row) else ""
            add_inline(p, value, size=9, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_NAVY)
    # Keep table rows from splitting across pages and repeat the header row.
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    header_tr_pr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc: Document, alt: str, rel_path: str) -> None:
    path = (ROOT / rel_path).resolve()
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_paragraph(doc, f"[Missing figure: {rel_path}]", center=True,
                      italic=True, color=GREY)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    run = caption.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.strip()) for c in cells)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def style_document(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Calibri")

    heading_specs = {
        "Heading 1": (16, NAVY, True, False),
        "Heading 2": (13, NAVY, True, False),
        "Heading 3": (11, NAVY, True, True),
    }
    for name, (size, color, bold, italic) in heading_specs.items():
        st = styles[name]
        st.font.name = "Cambria"
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = bold
        st.font.italic = italic
        st.paragraph_format.space_before = Pt(14)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Cambria"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = NAVY
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)


def build(
    md_path: Path = MD_PATH,
    out_path: Path = OUT_PATH,
    header: str = "Task-Oriented DeepJSCC for UAV Navigation",
    doc_title: str | None = None,
) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)
    enable_update_fields(doc)

    props = doc.core_properties
    props.title = doc_title or header
    props.subject = "Final Year Project Thesis"
    props.author = "Areesha Sohail; Amna Shahid; Hafiza Laiba Faisal"
    props.keywords = (
        "DeepJSCC; task-oriented communication; UAV navigation; AWGN; "
        "Rayleigh fading; 3GPP UMa-aerial; 6G"
    )
    props.comments = (
        f"Generated from {md_path.name} by tools/build_thesis_docx.py"
    )

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.different_first_page_header_footer = True
    add_header(section, header)
    add_page_number_footer(section)

    in_title = True
    first_heading = True
    pending_break = False
    suppress_next_break = False
    toc_added = False
    i = 0

    # A little breathing room above the title.
    doc.add_paragraph()
    doc.add_paragraph()

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped == PAGEBREAK:
            if suppress_next_break:
                suppress_next_break = False
            else:
                pending_break = True
            i += 1
            continue

        if stripped == TOC_MARKER:
            if not toc_added:
                add_toc(doc)
                toc_added = True
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Pipe table
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                if not is_separator_row(cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows, center=in_title)
            pending_break = False
            continue

        # Image
        m = IMAGE.match(stripped)
        if m:
            if pending_break:
                p = doc.add_paragraph()
                p.paragraph_format.page_break_before = True
                pending_break = False
            add_image(doc, m.group(1), m.group(2))
            i += 1
            continue

        # Horizontal rule: end of title page.
        if stripped == "---":
            if in_title:
                in_title = False
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            text = stripped[2:].strip()
            if first_heading:
                doc.add_heading(text, level=0)
                first_heading = False
            else:
                p = doc.add_heading(text, level=1)
                force_break = text.startswith("Appendix") or text == "References"
                if pending_break or force_break:
                    p.paragraph_format.page_break_before = True
                    pending_break = False
            i += 1
            continue

        if stripped.startswith("## "):
            text = stripped[3:].strip()
            p = doc.add_heading(text, level=2)
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
            i += 1
            continue

        if stripped.startswith("### "):
            text = stripped[4:].strip()
            p = doc.add_heading(text, level=3)
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
            i += 1
            continue

        # Bullets and numbered lists
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
            add_inline(p, stripped[2:].strip())
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        # Blockquote
        if stripped.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, stripped[2:].strip())
            i += 1
            continue

        # Ordinary paragraph
        if in_title:
            size = 14 if stripped.startswith("**Final Year Project Thesis") else 12
            add_paragraph(doc, stripped, center=True, size=size, space_after=8)
        else:
            p = doc.add_paragraph()
            if pending_break:
                p.paragraph_format.page_break_before = True
                pending_break = False
            add_inline(p, stripped)
        i += 1

    doc.save(out_path)
    print(f"Wrote {out_path} ({out_path.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--md", type=Path, default=MD_PATH)
    parser.add_argument("--out", type=Path, default=OUT_PATH)
    parser.add_argument(
        "--header",
        default="Task-Oriented DeepJSCC for UAV Navigation",
        help="Running header text",
    )
    parser.add_argument(
        "--title",
        default=None,
        help="Document title metadata (defaults to the header)",
    )
    args = parser.parse_args()
    build(args.md, args.out, args.header, args.title)
